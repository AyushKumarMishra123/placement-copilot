import io
import docx
from typing import Dict, Any, List
from ..models.resume_schema import VisualLayoutAnalysis

class DOCXParsingError(Exception):
    def __init__(self, message: str, code: str = "DOCX_PARSING_FAILURE"):
        super().__init__(message)
        self.message = message
        self.code = code

def parse_docx_document(file_bytes: bytes) -> Dict[str, Any]:
    """
    Extracts text, structural paragraphs, and tables from a DOCX document.
    """
    if not file_bytes or len(file_bytes) == 0:
        raise DOCXParsingError("Uploaded DOCX file is completely empty (0 bytes).", "EMPTY_DOCUMENT")

    try:
        doc_stream = io.BytesIO(file_bytes)
        doc = docx.Document(doc_stream)
    except Exception as e:
        raise DOCXParsingError(f"Failed to open or decode DOCX document: {str(e)}", "DOCX_PARSING_FAILURE")

    text_parts: List[str] = []

    # 1. Read standard paragraphs
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            text_parts.append(t)

    # 2. Read text inside tables (resumes often use tables for contact or columns)
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                # Deduplicate identical adjacent cells (due to cell merges)
                deduped = []
                for cell_t in row_cells:
                    if not deduped or deduped[-1] != cell_t:
                        deduped.append(cell_t)
                text_parts.append(" | ".join(deduped))

    combined_text = "\n".join(text_parts).strip()
    char_count = len(combined_text)
    words = combined_text.split()
    word_count = len(words)
    lines = [l for l in combined_text.split("\n") if l.strip()]
    line_count = len(lines)

    if char_count == 0:
        raise DOCXParsingError("DOCX document contains no text content.", "EMPTY_DOCUMENT")

    # Estimate page count: Standard single-spaced A4 resume is typically 450-600 words or 40-50 lines
    estimated_page_count = max(1, int(round(word_count / 520.0)))
    layout_anomalies: List[str] = []
    if estimated_page_count > 1:
        layout_anomalies.append(f"Estimated document length exceeds 1 page (approx {estimated_page_count} pages based on word count).")

    layout_analysis = VisualLayoutAnalysis(
        isScanned=False,
        confidence=0.95,
        layoutAnomalies=layout_anomalies,
        lineCount=line_count,
        wordCount=word_count,
        characterCount=char_count,
        estimatedSectionSpacing="normal"
    )

    return {
        "pageCount": estimated_page_count,
        "rawText": combined_text,
        "layout": layout_analysis
    }
